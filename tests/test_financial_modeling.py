import pandas as pd
import pytest
from openpyxl import load_workbook

from financial_modeling.data_loader import load_financial_data
from financial_modeling.validation import validate_financial_data
from financial_modeling.forecast import forecast_free_cash_flow
from financial_modeling.valuation import calculate_dcf
from financial_modeling.reporting import (
    calculate_sensitivity,
    export_valuation_report,
)
from financial_modeling.config import EXCEL_REPORT


def validated_data():
    return validate_financial_data(load_financial_data())


def test_dataset_loads():
    df = validated_data()
    assert len(df) == 2
    assert "revenue" in df.columns


def test_missing_column_rejected():
    df = validated_data().drop(columns=["revenue"])
    with pytest.raises(ValueError, match="Missing required columns"):
        validate_financial_data(df)


def test_duplicate_year_rejected():
    df = validated_data()
    df.loc[df.index[1], "year"] = df.loc[df.index[0], "year"]
    with pytest.raises(ValueError, match="duplicated"):
        validate_financial_data(df)


def test_forecast_has_five_years():
    forecast = forecast_free_cash_flow(validated_data())
    assert len(forecast) == 5
    assert forecast["year"].tolist() == [2027, 2028, 2029, 2030, 2031]
    assert (forecast["fcff"] > 0).all()


def test_dcf_components_add_up():
    forecast = forecast_free_cash_flow(validated_data())
    valuation = calculate_dcf(forecast)
    assert valuation["enterprise_value"] == pytest.approx(
        valuation["pv_forecast_cash_flows"]
        + valuation["terminal_present_value"]
    )
    assert valuation["enterprise_value"] > 0


def test_invalid_dcf_rate_rejected():
    forecast = forecast_free_cash_flow(validated_data())
    original = __import__(
        "financial_modeling.valuation", fromlist=["DISCOUNT_RATE"]
    ).DISCOUNT_RATE

    try:
        import financial_modeling.valuation as valuation_module
        valuation_module.DISCOUNT_RATE = 0.01
        valuation_module.TERMINAL_GROWTH_RATE = 0.025

        with pytest.raises(ValueError, match="Discount rate"):
            calculate_dcf(forecast)
    finally:
        valuation_module.DISCOUNT_RATE = original
        from financial_modeling.config import TERMINAL_GROWTH_RATE
        valuation_module.TERMINAL_GROWTH_RATE = TERMINAL_GROWTH_RATE


def test_sensitivity_has_nine_combinations():
    forecast = forecast_free_cash_flow(validated_data())
    sensitivity = calculate_sensitivity(forecast)
    assert len(sensitivity) == 9
    assert sensitivity["enterprise_value"].gt(0).all()


def test_excel_report_has_three_sheets():
    forecast = forecast_free_cash_flow(validated_data())
    valuation = calculate_dcf(forecast)
    sensitivity = calculate_sensitivity(forecast)

    export_valuation_report(forecast, valuation, sensitivity)

    workbook = load_workbook(EXCEL_REPORT, read_only=True)
    try:
        assert workbook.sheetnames == [
            "Valuation Summary",
            "Cash Flow Forecast",
            "Sensitivity Analysis",
        ]
    finally:
        workbook.close()